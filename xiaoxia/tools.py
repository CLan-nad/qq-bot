from __future__ import annotations

import ast
import asyncio
import html
import ipaddress
import math
import os
import re
import socket
import subprocess
from dataclasses import dataclass, field
from io import BytesIO
from typing import Any
from urllib.parse import urlparse

import httpx2


USER_AGENT = "Mozilla/5.0 (Windows NT 10.0; Win64; x64) xiaoxia-qq-bot/0.2"

# 服务器在境内，DuckDuckGo 系引擎基本连不上；Bing 网页端稳定且中文结果好，作为默认兜底。
BING_SEARCH_URL = "https://www.bing.com/search"
BING_SEARCH_FALLBACK_URL = "https://cn.bing.com/search"
BROWSER_USER_AGENT = (
    "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
    "(KHTML, like Gecko) Chrome/124.0.0.0 Safari/537.36"
)
_TAG_PATTERN = re.compile(r"<[^>]+>")


def _clean_html_text(value: str) -> str:
    text = _TAG_PATTERN.sub("", value)
    text = html.unescape(text).replace("\u2002", " ").replace("\xa0", " ")
    return re.sub(r"\s+", " ", text).strip()


@dataclass
class ToolResult:
    ok: bool
    output: str = ""
    data: list[dict[str, Any]] = field(default_factory=list)
    error: str = ""

    @property
    def text(self) -> str:
        return self.output or self.error


@dataclass(frozen=True)
class DocumentAttachment:
    url: str
    filename: str
    content_type: str


DOCUMENT_EXTENSIONS = frozenset(
    {".pdf", ".docx", ".doc", ".xlsx", ".xls", ".txt", ".md", ".markdown", ".csv"}
)
DOCUMENT_CONTENT_TYPES = frozenset(
    {
        "application/pdf",
        "application/x-pdf",
        "application/msword",
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        "application/vnd.ms-excel",
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        "text/plain",
        "text/markdown",
        "text/csv",
    }
)


def document_attachments(raw: dict[str, Any]) -> list[DocumentAttachment]:
    result: list[DocumentAttachment] = []
    for item in raw.get("attachments") or []:
        if not isinstance(item, dict):
            continue
        url = str(item.get("url") or "").strip()
        if not url.startswith(("http://", "https://")):
            continue
        filename = str(item.get("filename") or "")
        content_type = str(item.get("content_type") or "").lower().split(";", 1)[0]
        lower_name = filename.lower()
        if content_type not in DOCUMENT_CONTENT_TYPES and not any(
            lower_name.endswith(ext) for ext in DOCUMENT_EXTENSIONS
        ):
            continue
        result.append(DocumentAttachment(url=url, filename=filename, content_type=content_type))
    return result


_PRIVATE_NETWORKS = tuple(
    ipaddress.ip_network(network)
    for network in (
        "0.0.0.0/8",
        "10.0.0.0/8",
        "100.64.0.0/10",
        "127.0.0.0/8",
        "169.254.0.0/16",
        "172.16.0.0/12",
        "192.168.0.0/16",
        "224.0.0.0/4",
        "240.0.0.0/4",
        "::1/128",
        "fc00::/7",
        "fe80::/10",
        "::/128",
    )
)


def is_private_host(host: str) -> bool:
    try:
        addresses = socket.getaddrinfo(host, None)
    except OSError:
        return True
    for address in addresses:
        ip = ipaddress.ip_address(address[4][0])
        if any(ip in network for network in _PRIVATE_NETWORKS):
            return True
    return False


async def _download_bytes(
    url: str,
    *,
    max_bytes: int,
    timeout: float,
) -> tuple[bytes, str] | None:
    parsed = urlparse(url)
    if parsed.scheme not in {"http", "https"} or not parsed.hostname:
        return None
    if is_private_host(parsed.hostname):
        return None
    if max_bytes <= 0:
        return None
    try:
        async with httpx2.AsyncClient(follow_redirects=True, timeout=timeout) as client:
            response = await client.get(url, headers={"User-Agent": USER_AGENT})
            response.raise_for_status()
    except httpx2.HTTPError:
        return None
    data = response.content
    if not data or len(data) > max_bytes:
        return None
    content_type = response.headers.get("content-type", "").split(";", 1)[0].strip().lower()
    return data, content_type


async def search_web(
    query: str,
    *,
    limit: int = 5,
    provider: str = "auto",
    tavily_api_key: str = "",
    bocha_api_key: str = "",
    timeout: float = 12.0,
) -> ToolResult:
    query = query.strip()
    if not query:
        return ToolResult(ok=False, error="搜索词为空")

    provider = (provider or "auto").lower()
    if provider == "auto":
        if tavily_api_key:
            chain = ["tavily"]
        elif bocha_api_key:
            chain = ["bocha"]
        else:
            chain = ["bing_web"]
        if "bing_web" not in chain:
            chain.append("bing_web")
    else:
        chain = [provider]
        if provider != "bing_web":
            chain.append("bing_web")

    errors: list[str] = []
    for name in chain:
        if name == "tavily":
            if not tavily_api_key:
                errors.append("Tavily 未配置 key")
                continue
            result = await _tavily_search(query, limit, tavily_api_key, timeout)
        elif name == "bocha":
            if not bocha_api_key:
                errors.append("博查未配置 key")
                continue
            result = await _bocha_search(query, limit, bocha_api_key, timeout)
        elif name == "duckduckgo":
            result = await _duckduckgo_search(query, limit, timeout)
        elif name == "bing_web":
            result = await _bing_web_search(query, limit, timeout)
        else:
            errors.append(f"不支持的搜索提供方：{name}")
            continue

        if result.ok and (result.data or result.output):
            return result
        errors.append(result.error or f"{name} 没有返回结果")

    return ToolResult(ok=False, error="搜索失败：" + "；".join(errors))


async def _tavily_search(query: str, limit: int, api_key: str, timeout: float) -> ToolResult:
    try:
        async with httpx2.AsyncClient(timeout=timeout) as client:
            response = await client.post(
                "https://api.tavily.com/search",
                headers={"Authorization": f"Bearer {api_key}"},
                json={
                    "query": query,
                    "max_results": max(1, min(limit, 10)),
                    "search_depth": "basic",
                    "include_answer": True,
                },
            )
            response.raise_for_status()
            payload = response.json()
    except Exception as exc:
        return ToolResult(ok=False, error=f"Tavily 搜索失败：{exc}")

    data: list[dict[str, Any]] = []
    for item in payload.get("results") or []:
        data.append(
            {
                "title": item.get("title") or "",
                "url": item.get("url") or "",
                "snippet": item.get("content") or item.get("snippet") or "",
            }
        )
    return ToolResult(ok=True, output=str(payload.get("answer") or "").strip(), data=data)


async def _bocha_search(query: str, limit: int, api_key: str, timeout: float) -> ToolResult:
    """博查搜索：中文结果质量好，境内直连稳定。"""
    try:
        async with httpx2.AsyncClient(timeout=timeout) as client:
            response = await client.post(
                "https://api.bochaai.com/v1/web-search",
                headers={
                    "Authorization": f"Bearer {api_key}",
                    "Content-Type": "application/json",
                },
                json={
                    "query": query,
                    "count": max(1, min(limit, 10)),
                    "summary": True,
                },
            )
            response.raise_for_status()
            payload = response.json()
    except Exception as exc:
        return ToolResult(ok=False, error=f"博查搜索失败：{exc}")

    pages = ((payload.get("data") or {}).get("webPages") or {}).get("value") or []
    data = [
        {
            "title": item.get("name") or "",
            "url": item.get("url") or "",
            "snippet": item.get("summary") or item.get("snippet") or "",
        }
        for item in pages
        if isinstance(item, dict)
    ]
    if not data:
        return ToolResult(ok=False, error="博查没有返回可用结果")
    return ToolResult(ok=True, data=data)


async def _bing_web_search(query: str, limit: int, timeout: float) -> ToolResult:
    """抓取 Bing 网页结果，不需要 key，是境内服务器上最稳的兜底方案。"""
    errors: list[str] = []
    for url in (BING_SEARCH_URL, BING_SEARCH_FALLBACK_URL):
        try:
            async with httpx2.AsyncClient(timeout=timeout, follow_redirects=True) as client:
                response = await client.get(
                    url,
                    params={"q": query, "setlang": "zh-CN", "ensearch": "0"},
                    headers={
                        "User-Agent": BROWSER_USER_AGENT,
                        "Accept-Language": "zh-CN,zh;q=0.9",
                    },
                )
                response.raise_for_status()
                page = response.text
        except Exception as exc:
            errors.append(f"{url} 请求失败：{exc}")
            continue

        data = _parse_bing_results(page, limit)
        if data:
            return ToolResult(ok=True, data=data)
        errors.append(f"{url} 没有解析到结果")

    return ToolResult(ok=False, error="Bing 搜索失败：" + "；".join(errors))


def _parse_bing_results(page: str, limit: int) -> list[dict[str, str]]:
    data: list[dict[str, str]] = []
    for block in re.findall(r'<li class="[^"]*b_algo[^"]*".*?</li>', page, re.S):
        url_match = re.search(r'<h2>.*?<a[^>]+href="([^"]+)"', block, re.S)
        if not url_match:
            continue
        title_match = re.search(r"<h2>.*?<a[^>]*>(.*?)</a>", block, re.S)
        snippet_match = re.search(r"<p[^>]*>(.*?)</p>", block, re.S)
        data.append(
            {
                "title": _clean_html_text(title_match.group(1)) if title_match else "",
                "url": html.unescape(url_match.group(1)),
                "snippet": _clean_html_text(snippet_match.group(1)) if snippet_match else "",
            }
        )
        if len(data) >= limit:
            break
    return data


async def _duckduckgo_search(query: str, limit: int, timeout: float) -> ToolResult:
    """保留的备选方案；境内网络经常超时，正常情况不会走到这里。"""
    def _run() -> list[dict[str, str]]:
        from ddgs import DDGS

        with DDGS(timeout=max(3.0, min(timeout, 8.0))) as ddgs:
            return list(ddgs.text(query, max_results=max(1, min(limit, 10))))

    try:
        items = await asyncio.wait_for(asyncio.to_thread(_run), timeout=timeout + 2)
    except asyncio.TimeoutError:
        return ToolResult(ok=False, error="DuckDuckGo 搜索超时")
    except Exception as exc:
        return ToolResult(ok=False, error=f"DuckDuckGo 搜索失败：{exc}")

    data = [
        {
            "title": item.get("title") or "",
            "url": item.get("href") or item.get("url") or "",
            "snippet": item.get("body") or "",
        }
        for item in items
    ]
    if not data:
        return ToolResult(ok=False, error="DuckDuckGo 没有返回结果")
    return ToolResult(ok=True, data=data)


def _plain_text(data: bytes, content_type: str) -> str:
    if "html" in content_type:
        import html2text

        return html2text.html2text(data.decode("utf-8", errors="replace"))
    for encoding in ("utf-8", "gb18030", "latin-1"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def _pdf_text(data: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(BytesIO(data))
    parts: list[str] = []
    for page in reader.pages[:50]:
        parts.append(page.extract_text() or "")
    return "\n".join(parts)


def _docx_text(data: bytes) -> str:
    from docx import Document

    document = Document(BytesIO(data))
    return "\n".join(paragraph.text for paragraph in document.paragraphs)


def _xlsx_text(data: bytes) -> str:
    from openpyxl import load_workbook

    workbook = load_workbook(BytesIO(data), read_only=True, data_only=True)
    lines: list[str] = []
    for sheet in workbook.worksheets:
        for row in sheet.iter_rows(values_only=True):
            cells = ["" if cell is None else str(cell) for cell in row]
            if any(cells):
                lines.append("\t".join(cells))
            if len(lines) >= 2000:
                break
        if len(lines) >= 2000:
            break
    return "\n".join(lines)


def extract_document_text(
    data: bytes,
    *,
    filename: str = "",
    content_type: str = "",
    max_chars: int = 24000,
) -> str:
    name = (filename or "").lower()
    content_type = (content_type or "").lower()
    if name.endswith(".pdf") or content_type in {"application/pdf", "application/x-pdf"} or data[:5] == b"%PDF-":
        text = _pdf_text(data)
    elif name.endswith(".docx") or content_type == (
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    ):
        text = _docx_text(data)
    elif name.endswith((".xlsx", ".xls")) or content_type in {
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        "application/vnd.ms-excel",
    }:
        text = _xlsx_text(data)
    elif name.endswith(".doc") or content_type == "application/msword":
        return "[旧版 .doc 暂不支持，请另存为 .docx 或 PDF 后再发]"
    else:
        text = _plain_text(data, content_type)
    return text[:max_chars]


async def read_document(
    url: str,
    *,
    filename: str = "",
    max_chars: int = 24000,
    max_bytes: int = 20 * 1024 * 1024,
) -> ToolResult:
    downloaded = await _download_bytes(url, max_bytes=max_bytes, timeout=60)
    if downloaded is None:
        return ToolResult(ok=False, error="文档下载失败或超出大小限制")
    data, content_type = downloaded
    try:
        text = extract_document_text(
            data,
            filename=filename,
            content_type=content_type,
            max_chars=max_chars,
        )
    except Exception as exc:
        return ToolResult(ok=False, error=f"文档解析失败：{exc}")
    if not text.strip():
        return ToolResult(ok=False, error="没有从文档中提取到文字内容")
    return ToolResult(ok=True, output=text)


async def fetch_page(
    url: str,
    *,
    max_chars: int = 8000,
    max_bytes: int = 2 * 1024 * 1024,
) -> ToolResult:
    downloaded = await _download_bytes(url, max_bytes=max_bytes, timeout=20)
    if downloaded is None:
        return ToolResult(ok=False, error="网页抓取失败、超时或链接指向内网地址")
    data, content_type = downloaded
    try:
        text = _plain_text(data, content_type)
    except Exception as exc:
        return ToolResult(ok=False, error=f"网页解析失败：{exc}")
    return ToolResult(ok=True, output=text[:max_chars])


_SAFE_FUNCS: dict[str, Any] = {
    "abs": abs,
    "round": round,
    "min": min,
    "max": max,
    "sum": sum,
    "pow": pow,
    "int": int,
    "float": float,
    "sqrt": math.sqrt,
    "sin": math.sin,
    "cos": math.cos,
    "tan": math.tan,
    "log": math.log,
    "log10": math.log10,
    "exp": math.exp,
    "pi": math.pi,
    "e": math.e,
}

_ALLOWED_MATH_NODES = frozenset(
    {
        ast.Expression,
        ast.BinOp,
        ast.UnaryOp,
        ast.Constant,
        ast.Add,
        ast.Sub,
        ast.Mult,
        ast.Div,
        ast.FloorDiv,
        ast.Mod,
        ast.Pow,
        ast.USub,
        ast.UAdd,
        ast.Name,
        ast.Call,
    }
)


class _MathValidator(ast.NodeVisitor):
    def visit_Name(self, node: ast.Name) -> None:
        if node.id not in _SAFE_FUNCS:
            raise ValueError(f"不允许使用名称：{node.id}")

    def visit_Call(self, node: ast.Call) -> None:
        if not isinstance(node.func, ast.Name) or node.func.id not in _SAFE_FUNCS:
            raise ValueError("不允许调用该函数")
        for argument in node.args:
            self.visit(argument)
        for keyword in node.keywords:
            self.visit(keyword.value)

    def visit_Attribute(self, node: ast.Attribute) -> None:
        raise ValueError("不允许属性访问")

    def visit_Subscript(self, node: ast.Subscript) -> None:
        raise ValueError("不允许下标访问")

    def generic_visit(self, node: ast.AST) -> None:
        if type(node) not in _ALLOWED_MATH_NODES:
            raise ValueError(f"不允许的语法节点：{type(node).__name__}")
        ast.NodeVisitor.generic_visit(self, node)


def eval_math(expression: str) -> ToolResult:
    try:
        tree = ast.parse(expression, mode="eval")
    except SyntaxError as exc:
        return ToolResult(ok=False, error=f"表达式无法解析：{exc}")

    try:
        _MathValidator().visit(tree)
    except ValueError as exc:
        return ToolResult(ok=False, error=f"表达式不安全：{exc}")

    environment = {"__builtins__": {}}
    environment.update(_SAFE_FUNCS)
    try:
        value = eval(compile(tree, "<math>", "eval"), environment)
    except Exception as exc:
        return ToolResult(ok=False, error=f"计算失败：{exc}")

    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return ToolResult(ok=True, output=str(value))


READONLY_COMMANDS = frozenset(
    {
        "ls", "cat", "head", "tail", "wc", "du", "df", "free", "uptime", "date",
        "uname", "pwd", "whoami", "id", "echo", "printf", "grep", "find", "stat",
        "env", "sort", "uniq", "tr", "cut", "sed", "awk", "less", "more",
        "basename", "dirname", "which", "type", "hostname", "expr",
    }
)

DENY_PATTERNS = (
    r"\bsudo\b",
    r"\bsu\s+-",
    r"\bdd\b",
    r"\bmkfs\b",
    r"\bshutdown\b",
    r"\breboot\b",
    r"\bhalt\b",
    r"\bpoweroff\b",
    r":\(\)\s*\{\s*:\|:&\s*\};",
    r"(curl|wget)[^|\n]*\|\s*(ba|z)?sh\b",
    r">\s*/dev/(sd|nvme|mmcblk|hd|vd)",
    r"\bchmod\s+-R\s+777\s+/",
    r"\bchown\s+-R",
    r"\brm\s+-[a-z]*r[a-z]*\s+(/|~)",
)


def validate_shell_command(command: str, mode: str) -> tuple[bool, str]:
    command = command.strip()
    if not command:
        return False, "命令为空"

    lowered = command.lower()
    for pattern in DENY_PATTERNS:
        if re.search(pattern, lowered):
            return False, f"命令包含禁止的模式：{pattern}"

    if mode == "readonly":
        first = lowered.split()[0]
        if first not in READONLY_COMMANDS:
            return False, f"只读模式不允许命令：{first or command}"
        if re.search(r"(^|[|;&])\s*(curl|wget|git)\b", command):
            return False, "只读模式不允许网络命令"
        if ">" in command or ">>" in command:
            return False, "只读模式不允许写入重定向"
    return True, ""


def _limit_resources() -> None:
    if os.name != "posix":
        return
    try:
        import resource

        limit = 256 * 1024 * 1024
        resource.setrlimit(resource.RLIMIT_AS, (limit, limit))
        resource.setrlimit(resource.RLIMIT_CPU, (10, 10))
        resource.setrlimit(resource.RLIMIT_FSIZE, (32 * 1024 * 1024, 32 * 1024 * 1024))
        resource.setrlimit(resource.RLIMIT_NOFILE, (64, 64))
        resource.setrlimit(resource.RLIMIT_NPROC, (32, 32))
    except Exception:
        pass


def _run_subprocess(command: str, cwd: str, timeout: float) -> subprocess.CompletedProcess[str]:
    kwargs: dict[str, Any] = {
        "cwd": cwd,
        "stdout": subprocess.PIPE,
        "stderr": subprocess.STDOUT,
        "text": True,
        "timeout": timeout,
        "env": {
            "PATH": "/usr/local/sbin:/usr/local/bin:/usr/sbin:/usr/bin:/sbin:/bin",
            "HOME": cwd,
            "LANG": "C.UTF-8",
        },
    }
    if os.name == "posix":
        kwargs["preexec_fn"] = _limit_resources
    return subprocess.run(["/bin/bash", "-c", command], **kwargs, check=False)


def _safe_cwd(base_dir: str) -> str:
    path = os.path.abspath(os.path.expanduser(base_dir or os.getcwd()))
    os.makedirs(path, exist_ok=True)
    return path


async def run_shell(
    command: str,
    *,
    cwd: str,
    mode: str = "readonly",
    timeout: float = 10,
    max_output_chars: int = 50000,
) -> ToolResult:
    allowed, error = validate_shell_command(command, mode)
    if not allowed:
        return ToolResult(ok=False, error=error)

    workdir = _safe_cwd(cwd)
    try:
        process = await asyncio.to_thread(
            _run_subprocess,
            command,
            workdir,
            timeout,
        )
    except subprocess.TimeoutExpired:
        return ToolResult(ok=False, error=f"命令超过 {timeout:g} 秒被终止")
    except Exception as exc:
        return ToolResult(ok=False, error=f"命令执行失败：{exc}")

    output = (process.stdout or "").strip()
    if not output and process.returncode != 0:
        output = "(命令没有输出)"
    return ToolResult(
        ok=process.returncode == 0,
        output=output[:max_output_chars],
        error="" if process.returncode == 0 else f"命令退出码 {process.returncode}",
    )
