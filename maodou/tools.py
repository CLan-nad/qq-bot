from __future__ import annotations

import ast
import asyncio
import ipaddress
import math
import os
import re
import socket
import subprocess
from dataclasses import dataclass, field
from io import BytesIO
from typing import Any
from urllib.parse import urlparse

import httpx2


USER_AGENT = "Mozilla/5.0 (Windows NT 10.0; Win64; x64) maodou-qq-bot/0.2"


@dataclass
class ToolResult:
    ok: bool
    output: str = ""
    data: list[dict[str, Any]] = field(default_factory=list)
    error: str = ""

    @property
    def text(self) -> str:
        return self.output or self.error


@dataclass(frozen=True)
class DocumentAttachment:
    url: str
    filename: str
    content_type: str


DOCUMENT_EXTENSIONS = frozenset(
    {".pdf", ".docx", ".doc", ".xlsx", ".xls", ".txt", ".md", ".markdown", ".csv"}
)
DOCUMENT_CONTENT_TYPES = frozenset(
    {
        "application/pdf",
        "application/x-pdf",
        "application/msword",
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        "application/vnd.ms-excel",
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        "text/plain",
        "text/markdown",
        "text/csv",
    }
)


def document_attachments(raw: dict[str, Any]) -> list[DocumentAttachment]:
    result: list[DocumentAttachment] = []
    for item in raw.get("attachments") or []:
        if not isinstance(item, dict):
            continue
        url = str(item.get("url") or "").strip()
        if not url.startswith(("http://", "https://")):
            continue
        filename = str(item.get("filename") or "")
        content_type = str(item.get("content_type") or "").lower().split(";", 1)[0]
        lower_name = filename.lower()
        if content_type not in DOCUMENT_CONTENT_TYPES and not any(
            lower_name.endswith(ext) for ext in DOCUMENT_EXTENSIONS
        ):
            continue
        result.append(DocumentAttachment(url=url, filename=filename, content_type=content_type))
    return result


_PRIVATE_NETWORKS = tuple(
    ipaddress.ip_network(network)
    for network in (
        "0.0.0.0/8",
        "10.0.0.0/8",
        "100.64.0.0/10",
        "127.0.0.0/8",
        "169.254.0.0/16",
        "172.16.0.0/12",
        "192.168.0.0/16",
        "224.0.0.0/4",
        "240.0.0.0/4",
        "::1/128",
        "fc00::/7",
        "fe80::/10",
        "::/128",
    )
)


def is_private_host(host: str) -> bool:
    try:
        addresses = socket.getaddrinfo(host, None)
    except OSError:
        return True
    for address in addresses:
        ip = ipaddress.ip_address(address[4][0])
        if any(ip in network for network in _PRIVATE_NETWORKS):
            return True
    return False


async def _download_bytes(
    url: str,
    *,
    max_bytes: int,
    timeout: float,
) -> tuple[bytes, str] | None:
    parsed = urlparse(url)
    if parsed.scheme not in {"http", "https"} or not parsed.hostname:
        return None
    if is_private_host(parsed.hostname):
        return None
    if max_bytes <= 0:
        return None
    try:
        async with httpx2.AsyncClient(follow_redirects=True, timeout=timeout) as client:
            response = await client.get(url, headers={"User-Agent": USER_AGENT})
            response.raise_for_status()
    except httpx2.HTTPError:
        return None
    data = response.content
    if not data or len(data) > max_bytes:
        return None
    content_type = response.headers.get("content-type", "").split(";", 1)[0].strip().lower()
    return data, content_type


DEEPSEEK_SEARCH_PATH = "/v1/tools/search"
DEEPSEEK_FETCH_PATH = "/v1/tools/fetch"
DEEPSEEK_RETRY_DELAY_SECONDS = 0.25
_DEEPSEEK_RETRYABLE_STATUSES = frozenset({408, 429, 500, 502, 503, 504})


def _deepseek_tools_url(base_url: str, path: str) -> str:
    base = (base_url or "https://api.deepseek.com").strip().rstrip("/")
    if base.endswith("/v1"):
        return base + path[len("/v1") :]
    return base + path


def _deepseek_error_detail(response: httpx2.Response) -> str:
    try:
        payload = response.json()
    except Exception:
        return ""
    if not isinstance(payload, dict):
        return ""
    error = payload.get("error")
    if isinstance(error, dict):
        return str(error.get("message") or "").strip()[:120]
    return str(error or "").strip()[:120]


def _deepseek_error_message(response: httpx2.Response, *, action: str) -> str:
    status = response.status_code
    if status == 401:
        return "DeepSeek 鉴权失败，请检查服务器上的 LLM_API_KEY"
    if status == 402:
        return "DeepSeek 余额不足，请先充值"
    if status == 403:
        return "DeepSeek 拒绝了这次请求"
    if status == 404:
        return "DeepSeek 没有找到可用的网页正文"
    if status == 429:
        return "DeepSeek 请求过于频繁，请稍后再试"
    if status in {500, 502, 503, 504}:
        return f"DeepSeek {action}服务暂时不可用（HTTP {status}）"
    detail = _deepseek_error_detail(response)
    suffix = f"：{detail}" if detail else ""
    return f"DeepSeek {action}失败（HTTP {status}）{suffix}"


async def _post_deepseek_tool(
    *,
    api_key: str,
    base_url: str,
    path: str,
    payload: dict[str, Any],
    timeout: float,
    action: str,
) -> tuple[dict[str, Any] | None, str]:
    url = _deepseek_tools_url(base_url, path)
    headers = {
        "Authorization": f"Bearer {api_key}",
        "Content-Type": "application/json",
    }
    last_error = f"DeepSeek {action}失败"
    for attempt in range(2):
        try:
            async with httpx2.AsyncClient(timeout=timeout) as client:
                response = await client.post(url, headers=headers, json=payload)
        except httpx2.TimeoutException:
            last_error = f"DeepSeek {action}超时"
        except httpx2.HTTPError as exc:
            last_error = f"DeepSeek {action}网络错误：{exc}"
        except Exception as exc:
            last_error = f"DeepSeek {action}失败：{exc}"
        else:
            if response.status_code < 400:
                try:
                    body = response.json()
                except Exception:
                    return None, f"DeepSeek {action}返回了无法解析的数据"
                if not isinstance(body, dict):
                    return None, f"DeepSeek {action}返回格式不正确"
                return body, ""
            last_error = _deepseek_error_message(response, action=action)
            if response.status_code not in _DEEPSEEK_RETRYABLE_STATUSES:
                return None, last_error
        if attempt == 0:
            await asyncio.sleep(DEEPSEEK_RETRY_DELAY_SECONDS)
    return None, last_error


async def search_web(
    query: str,
    *,
    limit: int = 5,
    api_key: str = "",
    base_url: str = "https://api.deepseek.com",
    original_question: str = "",
    timeout: float = 12.0,
) -> ToolResult:
    query = query.strip()
    if not query:
        return ToolResult(ok=False, error="搜索词为空")
    if not api_key.strip():
        return ToolResult(ok=False, error="联网搜索没有配置：缺少 LLM_API_KEY")

    payload: dict[str, Any] = {
        "query": query,
        "count": max(1, min(limit, 50)),
        "need_content": False,
    }
    original_question = original_question.strip()[:128]
    if original_question:
        payload["ori_question"] = original_question

    body, error = await _post_deepseek_tool(
        api_key=api_key,
        base_url=base_url,
        path=DEEPSEEK_SEARCH_PATH,
        payload=payload,
        timeout=timeout,
        action="搜索",
    )
    if body is None:
        return ToolResult(ok=False, error=error)

    data: list[dict[str, Any]] = []
    for item in body.get("results") or []:
        if not isinstance(item, dict):
            continue
        url = str(item.get("url") or "").strip()
        if not url:
            continue
        summary = str(item.get("summary") or "").strip()
        snippet = str(item.get("snippet") or summary).strip()
        data.append(
            {
                "title": str(item.get("title") or "").strip(),
                "url": url,
                "snippet": snippet,
                "summary": summary,
                "site_name": str(item.get("site_name") or "").strip(),
                "date": str(item.get("date") or "").strip(),
                "authority": item.get("authority"),
                "rerank_score": item.get("rerank_score"),
            }
        )
    if not data:
        return ToolResult(ok=False, error="DeepSeek 搜索没有返回可用结果")
    return ToolResult(ok=True, data=data)


def _plain_text(data: bytes, content_type: str) -> str:
    if "html" in content_type:
        import html2text

        return html2text.html2text(data.decode("utf-8", errors="replace"))
    for encoding in ("utf-8", "gb18030", "latin-1"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def _pdf_text(data: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(BytesIO(data))
    parts: list[str] = []
    for page in reader.pages[:50]:
        parts.append(page.extract_text() or "")
    return "\n".join(parts)


def _docx_text(data: bytes) -> str:
    from docx import Document

    document = Document(BytesIO(data))
    return "\n".join(paragraph.text for paragraph in document.paragraphs)


def _xlsx_text(data: bytes) -> str:
    from openpyxl import load_workbook

    workbook = load_workbook(BytesIO(data), read_only=True, data_only=True)
    lines: list[str] = []
    for sheet in workbook.worksheets:
        for row in sheet.iter_rows(values_only=True):
            cells = ["" if cell is None else str(cell) for cell in row]
            if any(cells):
                lines.append("\t".join(cells))
            if len(lines) >= 2000:
                break
        if len(lines) >= 2000:
            break
    return "\n".join(lines)


def extract_document_text(
    data: bytes,
    *,
    filename: str = "",
    content_type: str = "",
    max_chars: int = 24000,
) -> str:
    name = (filename or "").lower()
    content_type = (content_type or "").lower()
    if name.endswith(".pdf") or content_type in {"application/pdf", "application/x-pdf"} or data[:5] == b"%PDF-":
        text = _pdf_text(data)
    elif name.endswith(".docx") or content_type == (
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    ):
        text = _docx_text(data)
    elif name.endswith((".xlsx", ".xls")) or content_type in {
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        "application/vnd.ms-excel",
    }:
        text = _xlsx_text(data)
    elif name.endswith(".doc") or content_type == "application/msword":
        return "[旧版 .doc 暂不支持，请另存为 .docx 或 PDF 后再发]"
    else:
        text = _plain_text(data, content_type)
    return text[:max_chars]


async def read_document(
    url: str,
    *,
    filename: str = "",
    max_chars: int = 24000,
    max_bytes: int = 20 * 1024 * 1024,
) -> ToolResult:
    downloaded = await _download_bytes(url, max_bytes=max_bytes, timeout=60)
    if downloaded is None:
        return ToolResult(ok=False, error="文档下载失败或超出大小限制")
    data, content_type = downloaded
    try:
        text = extract_document_text(
            data,
            filename=filename,
            content_type=content_type,
            max_chars=max_chars,
        )
    except Exception as exc:
        return ToolResult(ok=False, error=f"文档解析失败：{exc}")
    if not text.strip():
        return ToolResult(ok=False, error="没有从文档中提取到文字内容")
    return ToolResult(ok=True, output=text)


async def fetch_page(
    url: str,
    *,
    api_key: str = "",
    base_url: str = "https://api.deepseek.com",
    max_chars: int = 8000,
    timeout: float = 20.0,
) -> ToolResult:
    parsed = urlparse(url)
    if parsed.scheme not in {"http", "https"} or not parsed.hostname:
        return ToolResult(ok=False, error="链接格式不正确")
    if is_private_host(parsed.hostname):
        return ToolResult(ok=False, error="链接指向内网地址，不能读取")
    if not api_key.strip():
        return ToolResult(ok=False, error="网页抓取没有配置：缺少 LLM_API_KEY")

    body, error = await _post_deepseek_tool(
        api_key=api_key,
        base_url=base_url,
        path=DEEPSEEK_FETCH_PATH,
        payload={"url": url},
        timeout=timeout,
        action="网页抓取",
    )
    if body is None:
        return ToolResult(ok=False, error=error)

    content = str(body.get("content") or "").strip()
    if not content:
        return ToolResult(
            ok=False,
            error="DeepSeek 没有提取到网页正文，页面可能被目标站点拦截",
        )
    title = str(body.get("title") or "").strip()
    output = f"标题：{title}\n\n{content}" if title else content
    return ToolResult(ok=True, output=output[:max_chars])


_SAFE_FUNCS: dict[str, Any] = {
    "abs": abs,
    "round": round,
    "min": min,
    "max": max,
    "sum": sum,
    "pow": pow,
    "int": int,
    "float": float,
    "sqrt": math.sqrt,
    "sin": math.sin,
    "cos": math.cos,
    "tan": math.tan,
    "log": math.log,
    "log10": math.log10,
    "exp": math.exp,
    "pi": math.pi,
    "e": math.e,
}

_ALLOWED_MATH_NODES = frozenset(
    {
        ast.Expression,
        ast.BinOp,
        ast.UnaryOp,
        ast.Constant,
        ast.Add,
        ast.Sub,
        ast.Mult,
        ast.Div,
        ast.FloorDiv,
        ast.Mod,
        ast.Pow,
        ast.USub,
        ast.UAdd,
        ast.Name,
        ast.Call,
    }
)


class _MathValidator(ast.NodeVisitor):
    def visit_Name(self, node: ast.Name) -> None:
        if node.id not in _SAFE_FUNCS:
            raise ValueError(f"不允许使用名称：{node.id}")

    def visit_Call(self, node: ast.Call) -> None:
        if not isinstance(node.func, ast.Name) or node.func.id not in _SAFE_FUNCS:
            raise ValueError("不允许调用该函数")
        for argument in node.args:
            self.visit(argument)
        for keyword in node.keywords:
            self.visit(keyword.value)

    def visit_Attribute(self, node: ast.Attribute) -> None:
        raise ValueError("不允许属性访问")

    def visit_Subscript(self, node: ast.Subscript) -> None:
        raise ValueError("不允许下标访问")

    def generic_visit(self, node: ast.AST) -> None:
        if type(node) not in _ALLOWED_MATH_NODES:
            raise ValueError(f"不允许的语法节点：{type(node).__name__}")
        ast.NodeVisitor.generic_visit(self, node)


def eval_math(expression: str) -> ToolResult:
    try:
        tree = ast.parse(expression, mode="eval")
    except SyntaxError as exc:
        return ToolResult(ok=False, error=f"表达式无法解析：{exc}")

    try:
        _MathValidator().visit(tree)
    except ValueError as exc:
        return ToolResult(ok=False, error=f"表达式不安全：{exc}")

    environment = {"__builtins__": {}}
    environment.update(_SAFE_FUNCS)
    try:
        value = eval(compile(tree, "<math>", "eval"), environment)
    except Exception as exc:
        return ToolResult(ok=False, error=f"计算失败：{exc}")

    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return ToolResult(ok=True, output=str(value))


READONLY_COMMANDS = frozenset(
    {
        "ls", "cat", "head", "tail", "wc", "du", "df", "free", "uptime", "date",
        "uname", "pwd", "whoami", "id", "echo", "printf", "grep", "find", "stat",
        "env", "sort", "uniq", "tr", "cut", "sed", "awk", "less", "more",
        "basename", "dirname", "which", "type", "hostname", "expr",
    }
)

DENY_PATTERNS = (
    r"\bsudo\b",
    r"\bsu\s+-",
    r"\bdd\b",
    r"\bmkfs\b",
    r"\bshutdown\b",
    r"\breboot\b",
    r"\bhalt\b",
    r"\bpoweroff\b",
    r":\(\)\s*\{\s*:\|:&\s*\};",
    r"(curl|wget)[^|\n]*\|\s*(ba|z)?sh\b",
    r">\s*/dev/(sd|nvme|mmcblk|hd|vd)",
    r"\bchmod\s+-R\s+777\s+/",
    r"\bchown\s+-R",
    r"\brm\s+-[a-z]*r[a-z]*\s+(/|~)",
)


def validate_shell_command(command: str, mode: str) -> tuple[bool, str]:
    command = command.strip()
    if not command:
        return False, "命令为空"

    lowered = command.lower()
    for pattern in DENY_PATTERNS:
        if re.search(pattern, lowered):
            return False, f"命令包含禁止的模式：{pattern}"

    if mode == "readonly":
        first = lowered.split()[0]
        if first not in READONLY_COMMANDS:
            return False, f"只读模式不允许命令：{first or command}"
        if re.search(r"(^|[|;&])\s*(curl|wget|git)\b", command):
            return False, "只读模式不允许网络命令"
        if ">" in command or ">>" in command:
            return False, "只读模式不允许写入重定向"
    return True, ""


def _limit_resources() -> None:
    if os.name != "posix":
        return
    try:
        import resource

        limit = 256 * 1024 * 1024
        resource.setrlimit(resource.RLIMIT_AS, (limit, limit))
        resource.setrlimit(resource.RLIMIT_CPU, (10, 10))
        resource.setrlimit(resource.RLIMIT_FSIZE, (32 * 1024 * 1024, 32 * 1024 * 1024))
        resource.setrlimit(resource.RLIMIT_NOFILE, (64, 64))
        resource.setrlimit(resource.RLIMIT_NPROC, (32, 32))
    except Exception:
        pass


def _run_subprocess(command: str, cwd: str, timeout: float) -> subprocess.CompletedProcess[str]:
    kwargs: dict[str, Any] = {
        "cwd": cwd,
        "stdout": subprocess.PIPE,
        "stderr": subprocess.STDOUT,
        "text": True,
        "timeout": timeout,
        "env": {
            "PATH": "/usr/local/sbin:/usr/local/bin:/usr/sbin:/usr/bin:/sbin:/bin",
            "HOME": cwd,
            "LANG": "C.UTF-8",
        },
    }
    if os.name == "posix":
        kwargs["preexec_fn"] = _limit_resources
    return subprocess.run(["/bin/bash", "-c", command], **kwargs, check=False)


def _safe_cwd(base_dir: str) -> str:
    path = os.path.abspath(os.path.expanduser(base_dir or os.getcwd()))
    os.makedirs(path, exist_ok=True)
    return path


async def run_shell(
    command: str,
    *,
    cwd: str,
    mode: str = "readonly",
    timeout: float = 10,
    max_output_chars: int = 50000,
) -> ToolResult:
    allowed, error = validate_shell_command(command, mode)
    if not allowed:
        return ToolResult(ok=False, error=error)

    workdir = _safe_cwd(cwd)
    try:
        process = await asyncio.to_thread(
            _run_subprocess,
            command,
            workdir,
            timeout,
        )
    except subprocess.TimeoutExpired:
        return ToolResult(ok=False, error=f"命令超过 {timeout:g} 秒被终止")
    except Exception as exc:
        return ToolResult(ok=False, error=f"命令执行失败：{exc}")

    output = (process.stdout or "").strip()
    if not output and process.returncode != 0:
        output = "(命令没有输出)"
    return ToolResult(
        ok=process.returncode == 0,
        output=output[:max_output_chars],
        error="" if process.returncode == 0 else f"命令退出码 {process.returncode}",
    )
